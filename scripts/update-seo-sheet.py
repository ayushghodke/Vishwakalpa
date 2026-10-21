#!/usr/bin/env python3
"""
scripts/update-seo-sheet.py
Automates updating Vishwakalpa_30_Day_SEO_Content_Plan.xlsx
with task completion status, percentage, published URLs, and Friday KPI reviews.
"""

import sys
import os
import datetime
import argparse
import openpyxl

DEFAULT_SHEET_PATH = r"C:\Users\ayush\OneDrive\Desktop\Ayush\Client Project\Vishwakalpa\Data\Vishwakalpa_30_Day_SEO_Content_Plan.xlsx"
SITE_BASE_URL = "https://www.vishwakalpa.com"

def load_workbook_safe(path):
    if not os.path.exists(path):
        raise FileNotFoundError(f"Spreadsheet not found at: {path}")
    return openpyxl.load_workbook(path)

def update_task_row(ws, match_date=None, match_title_substr=None, status="Done", completion_pct=1.0, published_url="", notes=None):
    """
    Finds row matching date (datetime/date string) or title substring and updates columns:
    Col 11: Status (Not Started, In Progress, Done, Blocked)
    Col 12: % Complete (0.0 - 1.0)
    Col 13: Published URL
    Col 14: Date Completed
    Col 15: Notes / Blockers
    """
    updated = False
    today_str = datetime.date.today().strftime("%Y-%m-%d")

    for row_idx in range(5, ws.max_row + 1):
        cell_date = ws.cell(row=row_idx, column=1).value
        cell_title = str(ws.cell(row=row_idx, column=4).value or "")

        matched = False
        if match_date:
            if isinstance(cell_date, (datetime.datetime, datetime.date)):
                date_str = cell_date.strftime("%Y-%m-%d")
                if date_str == match_date:
                    matched = True
            elif str(cell_date).startswith(match_date):
                matched = True

        if not matched and match_title_substr:
            if match_title_substr.lower() in cell_title.lower():
                matched = True

        if matched:
            ws.cell(row=row_idx, column=11).value = status
            ws.cell(row=row_idx, column=12).value = float(completion_pct)
            if published_url:
                ws.cell(row=row_idx, column=13).value = published_url
            if status == "Done":
                ws.cell(row=row_idx, column=14).value = today_str
            if notes:
                ws.cell(row=row_idx, column=15).value = notes
            print(f"[OK] Updated Row {row_idx}: '{cell_title[:45]}...' -> Status={status}, Complete={completion_pct*100}%, URL={published_url}")
            updated = True

    return updated

def update_weekly_kpi(wb, week_ending_date="2026-10-16", blogs_published=4, case_studies_published=1, notes="Week 1 catch-up complete. All scheduled technical guides and case study live."):
    if "Weekly KPI Tracker" not in wb.sheetnames:
        print("[WARN] Sheet 'Weekly KPI Tracker' not found.")
        return False

    ws = wb["Weekly KPI Tracker"]
    for row_idx in range(4, ws.max_row + 1):
        cell_date = ws.cell(row=row_idx, column=1).value
        matched = False
        if isinstance(cell_date, (datetime.datetime, datetime.date)):
            date_str = cell_date.strftime("%Y-%m-%d")
            if date_str == week_ending_date:
                matched = True
        elif str(cell_date).startswith(week_ending_date):
            matched = True

        if matched:
            ws.cell(row=row_idx, column=2).value = blogs_published
            ws.cell(row=row_idx, column=3).value = case_studies_published
            ws.cell(row=row_idx, column=7).value = "None. All SEO checks passing (valid title, meta desc, schema, sitemap)."
            ws.cell(row=row_idx, column=8).value = "Week 2 Focus: Plastics recycling plant guide, factory master planning, material flow."
            print(f"[OK] Updated Weekly KPI Tracker for {week_ending_date}: Blogs={blogs_published}, CaseStudies={case_studies_published}")
            return True

    return False

def print_summary(wb):
    ws = wb[wb.sheetnames[0]]
    print(f"\n--- Current Schedule Status ({wb.sheetnames[0]}) ---")
    print(f"{'Date':<12} | {'Workstream':<16} | {'Status':<12} | {'%':<5} | {'Title':<45} | {'URL'}")
    print("-" * 110)
    for r in range(5, ws.max_row + 1):
        d = ws.cell(row=r, column=1).value
        d_str = d.strftime("%Y-%m-%d") if isinstance(d, (datetime.datetime, datetime.date)) else str(d)[:10]
        ws_type = str(ws.cell(row=r, column=3).value or "")[:15]
        status = str(ws.cell(row=r, column=11).value or "Not Started")
        pct = ws.cell(row=r, column=12).value
        pct_str = f"{int(pct * 100)}%" if pct is not None else "0%"
        title = str(ws.cell(row=r, column=4).value or "")[:43]
        url = str(ws.cell(row=r, column=13).value or "")
        print(f"{d_str:<12} | {ws_type:<16} | {status:<12} | {pct_str:<5} | {title:<45} | {url}")

def main():
    parser = argparse.ArgumentParser(description="Update Vishwakalpa SEO Schedule Excel Sheet")
    parser.add_argument("--sheet", default=DEFAULT_SHEET_PATH, help="Path to .xlsx file")
    parser.add_argument("--verify", action="store_true", help="Print summary of sheet without modifying")
    parser.add_argument("--week1-catchup", action="store_true", help="Update Week 1 (Oct 13 - Oct 16) tasks to completed")
    parser.add_argument("--day5", action="store_true", help="Update Day 5 (Oct 19) Plastic Recycling Plant task")
    parser.add_argument("--days6-7", action="store_true", help="Update Day 6 (Oct 20) and Day 7 (Oct 21) tasks")

    args = parser.parse_args()
    wb = load_workbook_safe(args.sheet)

    if args.verify:
        print_summary(wb)
        return

    ws_sched = wb[wb.sheetnames[0]]

    if args.week1_catchup:
        # Day 1: Oct 13 - What Is Industrial Architecture?
        update_task_row(ws_sched, match_date="2026-10-13", status="Done", completion_pct=1.0, 
                        published_url=f"{SITE_BASE_URL}/blogs/what-is-industrial-architecture",
                        notes="Published in-depth guide on industrial architecture vs commercial design.")

        # Day 2: Oct 14 - How to Design a Precision Machining Unit
        update_task_row(ws_sched, match_date="2026-10-14", status="Done", completion_pct=1.0, 
                        published_url=f"{SITE_BASE_URL}/blogs/how-to-design-a-precision-machining-unit",
                        notes="Published practical CNC machining plant guide with layouts, utilities & foundation details.")

        # Day 3: Oct 15 - How to Plan an Efficient Foundry Layout
        update_task_row(ws_sched, match_date="2026-10-15", status="Done", completion_pct=1.0, 
                        published_url=f"{SITE_BASE_URL}/blogs/how-to-plan-an-efficient-foundry-layout",
                        notes="Published sector guide on foundry layout, furnace zoning, crane clearance & ventilation.")

        # Day 4: Oct 16 - What to Check Before Expanding a Plant & Case Study 1
        update_task_row(ws_sched, match_date="2026-10-16", status="Done", completion_pct=1.0, 
                        published_url=f"{SITE_BASE_URL}/blogs/manufacturing-plant-expansion-planning",
                        notes="Published expansion planning guide. Verified Case Study 1 at /projects/precision-machining-facility-solapur.")

        # Update Friday KPI Review
        update_weekly_kpi(wb, week_ending_date="2026-10-16", blogs_published=4, case_studies_published=1)

    if args.day5:
        # Day 5: Oct 19 - Plastic Recycling Plant Design
        update_task_row(ws_sched, match_date="2026-10-19", status="Done", completion_pct=1.0,
                        published_url=f"{SITE_BASE_URL}/blogs/how-to-design-a-plastic-recycling-plant",
                        notes="Published guide on plastic recycling plant layout, wash lines, closed-loop ETP & NBC fire zoning.")

    if getattr(args, 'days6_7', False):
        # Day 6: Oct 20 - Factory Master Planning Decisions
        update_task_row(ws_sched, match_date="2026-10-20", status="Done", completion_pct=1.0,
                        published_url=f"{SITE_BASE_URL}/blogs/factory-master-planning-decisions",
                        notes="Published guide on 10 pre-construction master planning decisions, site logistics & expansion.")

        # Day 7: Oct 21 - Material Flow Planning
        update_task_row(ws_sched, match_date="2026-10-21", status="Done", completion_pct=1.0,
                        published_url=f"{SITE_BASE_URL}/blogs/how-to-plan-material-flow-manufacturing-facility",
                        notes="Published guide on material flow topologies (I/U/Cellular), WIP buffering & forklift safety.")

    wb.save(args.sheet)
    print("\n[SUCCESS] Saved updates to:", args.sheet)
    print_summary(wb)

if __name__ == "__main__":
    main()
